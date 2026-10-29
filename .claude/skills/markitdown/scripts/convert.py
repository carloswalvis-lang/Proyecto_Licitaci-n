#!/usr/bin/env python3
"""
MarkItDown converter — PDF, DOCX, PPTX -> Markdown
Uso: python convert.py <archivo_entrada> <archivo_salida.md> [--origen]
     --origen agrega como primera línea el nombre y la ruta del original (regla BD/ de docs/workflows.md).

1. Usa markitdown (instala markitdown[pdf,docx,pptx] si falta).
2. Si no se puede, usa un respaldo por formato:
   .docx -> python-docx | .pptx -> python-pptx | .pdf -> pdfplumber o pypdf

Códigos de salida: 0 OK | 1 uso | 2 no encontrado | 3 formato | 4 sin conversor | 5 error
"""
import os
import subprocess
import sys

SUPPORTED = {".pdf", ".docx", ".pptx"}
PKG = "markitdown[pdf,docx,pptx]"


def log(msg):
    print(msg, file=sys.stderr)


def pip_install(*pkgs):
    """Instala paquetes con pip. Devuelve True si tuvo éxito."""
    base = [sys.executable, "-m", "pip", "install", "-q", "--disable-pip-version-check",
            "--timeout", "20", "--retries", "1", *pkgs]
    for extra in ([], ["--break-system-packages"], ["--user"]):
        try:
            r = subprocess.run(base + extra, capture_output=True, text=True, timeout=180)
        except (subprocess.TimeoutExpired, OSError):
            return False
        if r.returncode == 0:
            return True
        err = (r.stderr or "").lower()
        if "externally-managed" not in err and "permission" not in err:
            # Error distinto (p. ej. sin red): no tiene sentido reintentar
            log(f"pip no pudo instalar {' '.join(pkgs)}: {(r.stderr or '').strip().splitlines()[-1:]}")
            return False
    return False


def get_markitdown():
    try:
        from markitdown import MarkItDown
        return MarkItDown
    except ImportError:
        log(f"MarkItDown no está instalado -> instalando {PKG}...")
        if pip_install(PKG):
            try:
                from markitdown import MarkItDown
                return MarkItDown
            except ImportError:
                pass
        return None


def is_missing_dependency(exc):
    name = type(exc).__name__
    text = str(exc).lower()
    return "MissingDependency" in name or "dependenc" in text or "pip install" in text


# ---------- Respaldos ----------

def table_to_md(rows):
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    clean = lambda c: (c or "").replace("\n", " ").replace("|", "\\|").strip()
    out = ["| " + " | ".join(clean(c) for c in rows[0]) + " |",
           "| " + " | ".join(["---"] * width) + " |"]
    out += ["| " + " | ".join(clean(c) for c in r) + " |" for r in rows[1:]]
    return out


def docx_fallback(path):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(path)
    headings = {"Title": "#", "Subtitle": "##"}
    lines = []
    # Recorre el cuerpo en orden, para mantener tablas en su posición
    for el in doc.element.body.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(el, doc)
            text = p.text.strip()
            if not text:
                lines.append("")
                continue
            style = p.style.name if p.style is not None else ""
            if style in headings:
                lines.append(f"{headings[style]} {text}")
            elif style.startswith("Heading"):
                lvl = "".join(ch for ch in style if ch.isdigit()) or "1"
                lines.append(f"{'#' * min(int(lvl), 6)} {text}")
            elif style.startswith("List"):
                lines.append(f"- {text}")
            else:
                lines.append(text)
        elif tag == "tbl":
            t = Table(el, doc)
            lines.append("")
            lines += table_to_md([[c.text for c in r.cells] for r in t.rows])
            lines.append("")
    return "\n".join(lines)


def pptx_fallback(path):
    from pptx import Presentation

    prs = Presentation(path)
    lines = []
    for i, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        lines.append(f"## Diapositiva {i}" + (f": {title}" if title else ""))
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        lines.append(f"{'  ' * para.level}- {t}")
            elif getattr(shape, "has_table", False) and shape.has_table:
                lines.append("")
                lines += table_to_md([[c.text for c in r.cells] for r in shape.table.rows])
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"\n> Notas: {notes}")
        lines.append("")
    return "\n".join(lines)


def pdf_fallback(path):
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            return "\n\n".join((p.extract_text() or "") for p in pdf.pages)
    except ImportError:
        from pypdf import PdfReader
        return "\n\n".join((p.extract_text() or "") for p in PdfReader(path).pages)


FALLBACKS = {".docx": docx_fallback, ".pptx": pptx_fallback, ".pdf": pdf_fallback}


def convert(path, ext):
    MarkItDown = get_markitdown()
    if MarkItDown is not None:
        try:
            return MarkItDown().convert(path).text_content, "markitdown"
        except Exception as e:
            if not is_missing_dependency(e):
                raise
            log(f"Faltan dependencias de MarkItDown para {ext} -> instalando {PKG}...")
            if pip_install(PKG):
                try:
                    return MarkItDown().convert(path).text_content, "markitdown"
                except Exception as e2:
                    if not is_missing_dependency(e2):
                        raise
    log(f"Usando respaldo para {ext}...")
    try:
        return FALLBACKS[ext](path), "respaldo"
    except ImportError as e:
        log(f"Error: no hay MarkItDown ni librería de respaldo disponible ({e}).")
        log("Probablemente el entorno no tiene acceso a PyPI (pypi.org / files.pythonhosted.org).")
        sys.exit(4)


def main():
    args = [a for a in sys.argv[1:] if a != "--origen"]
    origen = "--origen" in sys.argv[1:]
    if len(args) != 2:
        log("Uso: python convert.py <archivo_entrada> <archivo_salida.md> [--origen]")
        sys.exit(1)

    src, dst = args[0], args[1]
    if not os.path.isfile(src):
        log(f"Error: archivo no encontrado -> '{src}'")
        sys.exit(2)

    ext = os.path.splitext(src)[1].lower()
    if ext not in SUPPORTED:
        log(f"Error: formato '{ext}' no soportado. Formatos válidos: {', '.join(sorted(SUPPORTED))}")
        sys.exit(3)

    try:
        content, method = convert(src, ext)
    except SystemExit:
        raise
    except Exception as e:
        log(f"Error durante la conversión: {type(e).__name__}: {e}")
        sys.exit(5)

    content = content or ""
    if not content.strip():
        log("Advertencia: el contenido generado está vacío.")
        log("Si es un PDF escaneado (solo imágenes), MarkItDown no puede extraer texto sin OCR.")

    if origen:
        content = f"> Origen: {os.path.basename(src)} — {os.path.abspath(src)}\n\n" + content

    os.makedirs(os.path.dirname(os.path.abspath(dst)), exist_ok=True)
    with open(dst, "w", encoding="utf-8") as f:
        f.write(content)

    print(f"OK: convertido -> {dst}")
    print(f"Método: {method}")
    print(f"Caracteres: {len(content)}")


if __name__ == "__main__":
    main()
