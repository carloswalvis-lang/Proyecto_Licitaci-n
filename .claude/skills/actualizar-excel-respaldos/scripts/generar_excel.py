#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generar_excel.py - Excel de cotizaciones de un proyecto (hojas Listado y ECO) desde BD PRECIOS
y, opcionalmente, copia de respaldos.

Lee Cotizaciones.csv con herramientas/bd_precios/bd_precios.py del repositorio (misma carpeta
de datos: BD_PRECIOS_DIR o "bd_precios" de config/local.json). No escribe en BD PRECIOS ni en FSChile.

Uso:
  python3 generar_excel.py --proyecto "Nombre" --salida "ruta/Registro_Cotizaciones.xlsx"
                           [--todas] [--sobrescribir]
                           [--fschile "raiz FSChile" --respaldo "carpeta RESPALDO"]

  --proyecto      filtra por Proyecto_origen (sin distinguir mayusculas ni espacios extra).
  --todas         incluye tambien las filas reemplazadas (por defecto solo vigentes).
  --sobrescribir  reemplaza el Excel de salida si existe (por defecto se detiene).
  --fschile/--respaldo  copia a RESPALDO/ los originales de cada PROV (PROV-NNNN_<archivo>)
                  y escribe MERC-NNNN_fuentes.txt con las fuentes de cada MERC.
                  Nunca reemplaza archivos existentes en RESPALDO/.

Salida: JSON por stdout. Codigo distinto de 0 = error (mensaje en stderr).
Requiere openpyxl.
"""
import argparse, json, os, shutil, sys

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.abspath(os.path.join(AQUI, os.pardir, os.pardir, os.pardir, os.pardir))
sys.path.insert(0, os.path.join(RAIZ, "herramientas", "bd_precios"))
import bd_precios  # noqa: E402  (resuelve la carpeta de datos al importar)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.stderr.write("Falta openpyxl: pip install openpyxl\n")
    sys.exit(3)


def norm(s):
    return " ".join(str(s or "").strip().lower().split())


def numero(v):
    try:
        return int(str(v).strip())
    except ValueError:
        try:
            return float(str(v).replace(",", "."))
        except ValueError:
            return v


def filas_proyecto(proyecto, todas):
    filas = bd_precios.leer(bd_precios.F_COT, bd_precios.COLS_COT)
    sel = [f for f in filas if norm(f.get("Proyecto_origen")) == norm(proyecto)]
    if not todas:
        sel = [f for f in sel if norm(f.get("Estado")) == "vigente"]
    return sel


def escribir_excel(filas, salida):
    cols = bd_precios.COLS_COT
    wb = Workbook()
    ws = wb.active
    ws.title = "Listado"
    for c, h in enumerate(cols, 1):
        ws.cell(row=1, column=c, value=h).font = Font(bold=True)
    for r, f in enumerate(filas, 2):
        for c, h in enumerate(cols, 1):
            v = f.get(h, "")
            ws.cell(row=r, column=c, value=numero(v) if h == "Precio_neto_CLP" else v)
    ws.freeze_panes = "B2"
    for c, h in enumerate(cols, 1):
        largo = max([len(str(h))] + [len(str(f.get(h, ""))) for f in filas])
        ws.column_dimensions[get_column_letter(c)].width = min(max(10, largo + 2), 60)

    eco = wb.create_sheet("ECO")
    for r, h in enumerate(cols, 1):
        eco.cell(row=r, column=1, value=h).font = Font(bold=True)
        for c, f in enumerate(filas, 2):
            v = f.get(h, "")
            celda = eco.cell(row=r, column=c, value=numero(v) if h == "Precio_neto_CLP" else v)
            celda.alignment = Alignment(wrap_text=True, vertical="top")
    eco.freeze_panes = "B2"
    eco.column_dimensions["A"].width = 22
    for c in range(2, len(filas) + 2):
        eco.column_dimensions[get_column_letter(c)].width = 32

    os.makedirs(os.path.dirname(os.path.abspath(salida)), exist_ok=True)
    wb.save(salida)


def copiar_respaldos(filas, fschile, respaldo):
    os.makedirs(respaldo, exist_ok=True)
    res = {"copiados": [], "ya_existian": [], "no_encontrados": [], "fuentes_merc": []}
    for f in filas:
        id_ = f.get("ID", "")
        if f.get("Tipo") == "PROV":
            rel = (f.get("Link_o_Ruta_FS") or "").strip().replace("\\", "/")
            origen = os.path.join(fschile, *[p for p in rel.split("/") if p])
            destino = os.path.join(respaldo, "%s_%s" % (id_, os.path.basename(rel)))
            if not rel or not os.path.isfile(origen):
                res["no_encontrados"].append({"ID": id_, "ruta": rel})
            elif os.path.exists(destino):
                res["ya_existian"].append(os.path.basename(destino))
            else:
                shutil.copy2(origen, destino)
                res["copiados"].append(os.path.basename(destino))
        elif f.get("Tipo") == "MERC":
            destino = os.path.join(respaldo, "%s_fuentes.txt" % id_)
            if os.path.exists(destino):
                res["ya_existian"].append(os.path.basename(destino))
                continue
            texto = "\n".join([
                "%s - %s" % (id_, f.get("Descripcion", "")),
                "Fecha cotizacion: %s" % f.get("Fecha_cotizacion", ""),
                "Precio neto CLP: %s / %s" % (f.get("Precio_neto_CLP", ""), f.get("Unidad", "")),
                "Proveedores: %s" % f.get("Proveedor_Fuentes", ""),
                "Link: %s" % f.get("Link_o_Ruta_FS", ""),
                "Fuentes:",
            ] + ["  - " + p.strip() for p in (f.get("Ubicacion_Detalle") or "").split(";") if p.strip()])
            with open(destino, "w", encoding="utf-8") as fh:
                fh.write(texto + "\n")
            res["fuentes_merc"].append(os.path.basename(destino))
    return res


def main():
    ap = argparse.ArgumentParser(description="Excel de cotizaciones del proyecto desde BD PRECIOS")
    ap.add_argument("--proyecto", required=True)
    ap.add_argument("--salida", required=True)
    ap.add_argument("--todas", action="store_true")
    ap.add_argument("--sobrescribir", action="store_true")
    ap.add_argument("--fschile")
    ap.add_argument("--respaldo")
    a = ap.parse_args()

    if bool(a.fschile) != bool(a.respaldo):
        sys.stderr.write("--fschile y --respaldo van juntos\n")
        sys.exit(1)
    if os.path.exists(a.salida) and not a.sobrescribir:
        sys.stderr.write("Ya existe %s; usa --sobrescribir solo con confirmacion del usuario\n" % a.salida)
        sys.exit(4)

    filas = filas_proyecto(a.proyecto, a.todas)
    if not filas:
        sys.stderr.write("No hay cotizaciones con Proyecto_origen = %r en BD PRECIOS\n" % a.proyecto)
        sys.exit(5)

    escribir_excel(filas, a.salida)
    out = {"excel": os.path.abspath(a.salida), "filas": len(filas),
           "PROV": sum(1 for f in filas if f.get("Tipo") == "PROV"),
           "MERC": sum(1 for f in filas if f.get("Tipo") == "MERC")}
    if a.fschile:
        if not os.path.isdir(a.fschile):
            sys.stderr.write("No existe la carpeta FSChile: %s\n" % a.fschile)
            sys.exit(6)
        out["respaldos"] = copiar_respaldos(filas, a.fschile, a.respaldo)
    print(json.dumps(out, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
