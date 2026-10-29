#!/usr/bin/env python3
# Escribe las filas de la nómina técnica en el Excel destino (solo agrega, nunca borra).
# Uso: python3 write_nomina.py "ruta/Nomina.xlsx" "Nomina" "rows.json"
import json, sys
from openpyxl import load_workbook

HEADERS = [
    "NOMBRE COMPLETO", "CARGO A DESEMPEÑAR", "TÍTULO PROFESIONAL", "A", "B",
    "CURSOS DE ESPECIALIZACIÓN", "EXPERIENCIA GENERAL (RESUMEN)",
    "EXPERIENCIA ESPEC (RESUMEN)", "OTRAS EXPERIENCIA GENERAL (RESUMEN)",
    "ANTECEDENTES ACADÉMICOS", "CUMPLE BBTT",
]

def norm(s):
    return " ".join(str(s or "").strip().lower().split())

def run(excel_path, sheet_name, rows_json_path):
    with open(rows_json_path, encoding="utf-8") as f:
        rows = json.load(f)  # list of dicts keyed by HEADERS

    wb = load_workbook(excel_path)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.create_sheet(sheet_name)

    # IMPORTANT: decide whether the sheet is empty before touching any cell.
    # Reading a cell (even ws[1] or ws.cell(1,1).value) makes openpyxl treat that
    # row as "used", which then shifts ws.append() to row 2 and leaves row 1 blank.
    # Writing every row explicitly with ws.cell(row=..., column=..., value=...)
    # instead of ws.append() avoids that trap entirely.
    is_empty = ws.max_row == 1 and ws.max_column == 1 and ws["A1"].value is None
    existing_names = set()

    if is_empty:
        for col, h in enumerate(HEADERS, start=1):
            ws.cell(row=1, column=col, value=h)
        next_row = 2
    else:
        first_row_values = [ws.cell(row=1, column=c).value for c in range(1, len(HEADERS) + 1)]
        if norm(first_row_values[0]) == norm(HEADERS[0]):
            next_row = ws.max_row + 1
            for r in range(2, ws.max_row + 1):
                existing_names.add(norm(ws.cell(row=r, column=1).value))
        else:
            # Sheet has other content but not our headers yet: add them after it.
            next_row = ws.max_row + 1
            for col, h in enumerate(HEADERS, start=1):
                ws.cell(row=next_row, column=col, value=h)
            next_row += 1

    added, skipped = 0, []
    for row in rows:
        name = norm(row.get("NOMBRE COMPLETO"))
        if name in existing_names:
            skipped.append(row.get("NOMBRE COMPLETO"))
            continue
        for col, h in enumerate(HEADERS, start=1):
            ws.cell(row=next_row, column=col, value=row.get(h, ""))
        existing_names.add(name)
        next_row += 1
        added += 1

    wb.save(excel_path)
    print(json.dumps({"added": added, "skipped_duplicates": skipped}, ensure_ascii=False))

if __name__ == "__main__":
    run(sys.argv[1], sys.argv[2], sys.argv[3])
